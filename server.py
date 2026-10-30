"""DocuFlow AI: Telegram Mini App + payments + document generation.
Run: uvicorn server:app --host 0.0.0.0 --port 8000
"""
import os, json, time, hmac, hashlib, sqlite3, secrets, re, io, asyncio, logging
from pathlib import Path
from urllib.parse import parse_qsl
from datetime import datetime, timezone
from contextlib import contextmanager
from xml.sax.saxutils import escape
from typing import Optional
import httpx
from fastapi import FastAPI, Request, HTTPException, UploadFile, File, Depends
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4, landscape
from docx import Document
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from PIL import Image as PILImage

ROOT=Path(__file__).resolve().parent
for line in (ROOT/'.env').read_text(encoding='utf-8').splitlines() if (ROOT/'.env').exists() else []:
    if '=' in line and not line.lstrip().startswith('#'):
        k,v=line.split('=',1); os.environ.setdefault(k.strip(),v.strip())
BOT_TOKEN=os.getenv('BOT_TOKEN',''); BOT_USERNAME=os.getenv('BOT_USERNAME','YourBot')
ADMIN_IDS={int(x) for x in os.getenv('ADMIN_IDS','8959461245').split(',') if x.strip().isdigit()}
PUBLIC_URL=os.getenv('PUBLIC_URL','').rstrip('/'); API_KEY=os.getenv('AI_API_KEY','')
AI_BASE=os.getenv('AI_BASE_URL','https://api.openai.com/v1').rstrip('/')
AI_MODEL=os.getenv('AI_MODEL','gpt-4.1-mini')
DB=ROOT/os.getenv('DB_PATH','data/docuflow.db'); DB.parent.mkdir(parents=True,exist_ok=True)
OUTPUT=ROOT/'data'/'outputs'; OUTPUT.mkdir(parents=True,exist_ok=True)
UPLOADS=ROOT/'data'/'uploads'; UPLOADS.mkdir(parents=True,exist_ok=True)
app=FastAPI(title='DocuFlow AI')
logging.basicConfig(level=logging.INFO)

@contextmanager
def conn():
    db=sqlite3.connect(DB,timeout=30);db.row_factory=sqlite3.Row
    try: yield db;db.commit()
    finally: db.close()

with conn() as db:
    db.executescript('''CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, name TEXT, free_left INTEGER DEFAULT 1, bonus_left INTEGER DEFAULT 0, premium_until INTEGER DEFAULT 0, referred_by INTEGER, created INTEGER);
    CREATE TABLE IF NOT EXISTS rewards(user_id INTEGER, milestone INTEGER, PRIMARY KEY(user_id,milestone));
    CREATE TABLE IF NOT EXISTS payments(charge_id TEXT PRIMARY KEY, user_id INTEGER, plan INTEGER, created INTEGER);
    CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY,user_id INTEGER,title TEXT,kind TEXT,config TEXT,content TEXT,created INTEGER);
    CREATE TABLE IF NOT EXISTS files(id TEXT PRIMARY KEY,project_id TEXT,user_id INTEGER,format TEXT,path TEXT,created INTEGER);
    CREATE TABLE IF NOT EXISTS jobs(id TEXT PRIMARY KEY,user_id INTEGER,status TEXT,error TEXT,project_id TEXT,created INTEGER);
    CREATE TABLE IF NOT EXISTS uploads(id TEXT PRIMARY KEY,user_id INTEGER,path TEXT,created INTEGER);
    ''')

def verify_init(raw):
    if not BOT_TOKEN: raise HTTPException(503,'BOT_TOKEN не настроен')
    if not raw: raise HTTPException(401,'Откройте приложение из Telegram')
    pairs=dict(parse_qsl(raw,keep_blank_values=True)); digest=pairs.pop('hash','')
    auth=int(pairs.get('auth_date','0'))
    if abs(time.time()-auth)>86400: raise HTTPException(401,'Авторизация устарела')
    check='\n'.join(f'{k}={v}' for k,v in sorted(pairs.items()))
    secret=hmac.new(b'WebAppData',BOT_TOKEN.encode(),hashlib.sha256).digest()
    expected=hmac.new(secret,check.encode(),hashlib.sha256).hexdigest()
    if not hmac.compare_digest(digest,expected):raise HTTPException(401,'Неверная подпись Telegram')
    try:return json.loads(pairs['user'])
    except Exception:raise HTTPException(401,'Нет пользователя Telegram')

def auth(req:Request):
    user=verify_init(req.headers.get('X-Telegram-Init-Data',''))
    uid=int(user['id']); name=user.get('first_name','Пользователь')
    with conn() as db:
        db.execute('INSERT OR IGNORE INTO users(id,name,created) VALUES(?,?,?)',(uid,name,int(time.time())))
        db.execute('UPDATE users SET name=? WHERE id=?',(name,uid))
    return uid

def reward(db,referrer):
    count=db.execute('SELECT COUNT(*) n FROM users WHERE referred_by=?',(referrer,)).fetchone()['n']
    for milestone in (5,10):
        if count>=milestone:
            inserted=db.execute('INSERT OR IGNORE INTO rewards(user_id,milestone) VALUES(?,?)',(referrer,milestone)).rowcount
            if inserted:
                if milestone==5:db.execute('UPDATE users SET bonus_left=bonus_left+3 WHERE id=?',(referrer,))
                if milestone==10:db.execute('UPDATE users SET premium_until=MAX(premium_until,?)+? WHERE id=?',(int(time.time()),7*86400,referrer))

@app.get('/')
def index():return FileResponse(ROOT/'index.html')

def admin_auth(uid=Depends(auth)):
    if uid not in ADMIN_IDS:raise HTTPException(403,'Доступ только администратору')
    return uid

@app.get('/api/admin/summary')
def admin_summary(uid=Depends(admin_auth)):
    with conn() as db:
        users=db.execute('SELECT COUNT(*) n FROM users').fetchone()['n']
        premium=db.execute('SELECT COUNT(*) n FROM users WHERE premium_until>?',(int(time.time()),)).fetchone()['n']
        projects=db.execute('SELECT COUNT(*) n FROM projects').fetchone()['n']
        payments=db.execute('SELECT COUNT(*) n FROM payments').fetchone()['n']
    return {'users':users,'premium':premium,'projects':projects,'payments':payments}

@app.get('/api/me')
def me(uid=Depends(auth)):
    with conn() as db:
        u=dict(db.execute('SELECT * FROM users WHERE id=?',(uid,)).fetchone())
        count=db.execute('SELECT COUNT(*) n FROM users WHERE referred_by=?',(uid,)).fetchone()['n']
        docs=db.execute('SELECT COUNT(*) n FROM projects WHERE user_id=?',(uid,)).fetchone()['n']
    return dict(id=uid,name=u['name'],free_left=u['free_left'],bonus_left=u['bonus_left'],premium_until=u['premium_until'],referrals=count,documents=docs,ref_link=f'https://t.me/{BOT_USERNAME}?startapp=ref_{uid}')

class Referral(BaseModel): code:str
@app.post('/api/referral')
def referral(data:Referral,uid=Depends(auth)):
    m=re.fullmatch(r'ref_(\d+)',data.code)
    if not m:return {'ok':False}
    ref=int(m.group(1))
    if ref==uid:return {'ok':False}
    with conn() as db:
        if not db.execute('SELECT id FROM users WHERE id=?',(ref,)).fetchone():return {'ok':False}
        row=db.execute('SELECT referred_by FROM users WHERE id=?',(uid,)).fetchone()
        if row['referred_by'] is not None:return {'ok':False}
        db.execute('UPDATE users SET referred_by=? WHERE id=?',(ref,uid));reward(db,ref)
    return {'ok':True}

PLANS={7:50,15:150,30:250}
class Plan(BaseModel):days:int
@app.post('/api/invoice')
async def invoice(data:Plan,uid=Depends(auth)):
    if data.days not in PLANS:raise HTTPException(400,'Неверный тариф')
    if not BOT_TOKEN:raise HTTPException(503,'Бот не настроен')
    payload=f'docuflow:{uid}:{data.days}:{secrets.token_hex(8)}'
    async with httpx.AsyncClient(timeout=20) as client:
        r=await client.post(f'https://api.telegram.org/bot{BOT_TOKEN}/createInvoiceLink',json={'title':f'DocuFlow Premium {data.days} дней','description':'Доступ к созданию документов','payload':payload,'currency':'XTR','prices':[{'label':'Premium','amount':PLANS[data.days]}]})
    response=r.json()
    if not response.get('ok'):raise HTTPException(502,'Telegram не создал счёт')
    return {'url':response['result']}

@app.post('/telegram/webhook')
async def telegram_webhook(req:Request):
    # Set webhook with secret_token via setup_webhook.py
    expected=os.getenv('WEBHOOK_SECRET','')
    if not expected or not hmac.compare_digest(req.headers.get('X-Telegram-Bot-Api-Secret-Token',''),expected):raise HTTPException(403)
    update=await req.json(); pre=update.get('pre_checkout_query')
    if pre:
        p=pre.get('invoice_payload','').split(':');ok=len(p)==4 and p[0]=='docuflow' and p[1]==str(pre.get('from',{}).get('id')) and p[2].isdigit() and int(p[2]) in PLANS and pre.get('currency')=='XTR' and pre.get('total_amount')==PLANS[int(p[2])]
        await tg_call('answerPreCheckoutQuery',{'pre_checkout_query_id':pre['id'],'ok':ok,'error_message':'Неверный платёж' if not ok else None})
    msg=update.get('message',{});payment=msg.get('successful_payment')
    if payment:
        p=payment.get('invoice_payload','').split(':');uid=msg['from']['id']
        if len(p)==4 and p[0]=='docuflow' and p[1]==str(uid) and p[2].isdigit() and int(p[2]) in PLANS and payment.get('currency')=='XTR' and payment.get('total_amount')==PLANS[int(p[2])]:
            days=int(p[2]);charge=payment['telegram_payment_charge_id']
            with conn() as db:
                db.execute('INSERT OR IGNORE INTO users(id,name,created) VALUES(?,?,?)',(uid,msg['from'].get('first_name',''),int(time.time())))
                if db.execute('INSERT OR IGNORE INTO payments VALUES(?,?,?,?)',(charge,uid,days,int(time.time()))).rowcount:
                    db.execute('UPDATE users SET premium_until=MAX(premium_until,?)+? WHERE id=?',(int(time.time()),days*86400,uid))
            await tg_call('sendMessage',{'chat_id':uid,'text':f'⭐ Оплата получена! Premium активирован на {days} дней.'})
    if msg.get('text','').startswith('/start'):
        await tg_call('sendMessage',{'chat_id':msg['chat']['id'],'text':'💜 Добро пожаловать в DocuFlow AI! Создавайте презентации и документы.','reply_markup':{'inline_keyboard':[[{'text':'🚀 Открыть DocuFlow AI','web_app':{'url':PUBLIC_URL}}]]} if PUBLIC_URL.startswith('https://') else None})
    return {'ok':True}

async def tg_call(method,payload):
    if not BOT_TOKEN:return
    async with httpx.AsyncClient(timeout=45) as client:
        try:
            r=await client.post(f'https://api.telegram.org/bot{BOT_TOKEN}/{method}',json=payload);r.raise_for_status();return r.json()
        except Exception:logging.exception('Telegram API %s failed',method)

@app.post('/api/upload')
async def upload(file:UploadFile=File(...),uid=Depends(auth)):
    raw=await file.read(8*1024*1024+1)
    if len(raw)>8*1024*1024:raise HTTPException(413,'Максимум 8 МБ')
    try:
        img=PILImage.open(io.BytesIO(raw));img.verify()
        img=PILImage.open(io.BytesIO(raw)).convert('RGB');img.thumbnail((1800,1800))
    except Exception:raise HTTPException(400,'Неверный формат изображения')
    fid=secrets.token_hex(12);path=UPLOADS/f'{fid}.jpg';img.save(path,quality=86)
    with conn() as db:db.execute('INSERT INTO uploads VALUES(?,?,?,?)',(fid,uid,str(path),int(time.time())))
    return {'id':fid}

class Config(BaseModel):
    topic:str=Field(min_length=3,max_length=300)
    prompt:str=Field(default='',max_length=3000)
    kind:str='Доклад'
    pages:int=Field(default=5,ge=1,le=30)
    level:str='Школа'
    theme:str='violet'
    background:str='plain'
    image_mode:str='none'
    images:list[Optional[str]]=Field(default_factory=list,max_length=30)
    formats:list[str]=Field(default_factory=lambda:['pdf'])

FONT='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
if not Path(FONT).exists():FONT='/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'
if Path(FONT).exists():pdfmetrics.registerFont(TTFont('DocuUnicode',FONT))
COLORS={'violet':'#7955FF','blue':'#328CFF','emerald':'#10A98B','rose':'#E05B91','amber':'#D68B22'}

def sanitize_html(s):return escape(str(s))

async def generate_ai(cfg:Config):
    if not API_KEY:raise RuntimeError('AI_API_KEY не настроен. Добавьте ключ в .env')
    instruction='''Ты автор качественных учебных материалов. Верни ТОЛЬКО валидный JSON объект: {"title":"...","slides":[{"heading":"...","paragraphs":["...","..."],"image_query":"English concise image query"}],"sources":["название или URL"]}. Количество slides строго равно заданному. Пиши на русском грамотным естественным языком. Не выдумывай библиографические источники и цитаты: если нет проверенных ссылок, sources пустой. Каждый слайд/страница имеет 1-3 содержательных абзаца, без повторов. Для презентации делай лаконично, для эссе связный аргументированный текст. Не включай markdown fences.'''
    request=f'Тип: {cfg.kind}. Тема: {cfg.topic}. Уровень: {cfg.level}. Страниц/слайдов: {cfg.pages}. Дополнительные пожелания: {cfg.prompt}'
    async with httpx.AsyncClient(timeout=115) as client:
        r=await client.post(f'{AI_BASE}/chat/completions',headers={'Authorization':f'Bearer {API_KEY}'},json={'model':AI_MODEL,'temperature':0.65,'response_format':{'type':'json_object'},'messages':[{'role':'system','content':instruction},{'role':'user','content':request}]})
        r.raise_for_status();answer=r.json()['choices'][0]['message']['content']
    obj=json.loads(answer)
    if not isinstance(obj.get('slides'),list) or not obj['slides']:raise RuntimeError('AI вернул пустой документ')
    obj['slides']=obj['slides'][:cfg.pages]
    return obj

async def image_auto(query):
    if not query:return None
    try:
        async with httpx.AsyncClient(timeout=12,headers={'User-Agent':'DocuFlowAI/1.0 (educational project)'}) as c:
            r=await c.get('https://commons.wikimedia.org/w/api.php',params={'action':'query','generator':'search','gsrsearch':query,'gsrnamespace':6,'gsrlimit':5,'prop':'imageinfo','iiprop':'url|mime','iiurlwidth':1000,'format':'json'});r.raise_for_status()
            for page in r.json().get('query',{}).get('pages',{}).values():
                info=(page.get('imageinfo') or [{}])[0]
                if info.get('mime') in ('image/jpeg','image/png'):
                    url=info.get('thumburl') or info.get('url')
                    resp=await c.get(url)
                    if resp.status_code==200 and len(resp.content)<5_000_000:
                        im=PILImage.open(io.BytesIO(resp.content)).convert('RGB');im.thumbnail((1300,900))
                        output=io.BytesIO();im.save(output,format='JPEG',quality=80);return output.getvalue()
    except Exception:logging.exception('Image search failed')
    return None

def add_pdf(content,cfg,imgs,path):
    accent=HexColor(COLORS.get(cfg.theme,COLORS['violet']))
    page_size=landscape(A4) if cfg.kind=='Презентация' else A4
    doc=SimpleDocTemplate(str(path),pagesize=page_size,rightMargin=48,leftMargin=48,topMargin=48,bottomMargin=48)
    style=ParagraphStyle('body',fontName='DocuUnicode',fontSize=11,leading=17,spaceAfter=13,textColor=HexColor('#242236'))
    heading=ParagraphStyle('heading',parent=style,fontSize=19,leading=25,textColor=accent,spaceAfter=18)
    cover=ParagraphStyle('cover',parent=style,fontSize=25,leading=35,alignment=TA_CENTER,textColor=accent)
    story=[Spacer(1,100),Paragraph(sanitize_html(content.get('title',cfg.topic)),cover),Spacer(1,30),Paragraph(sanitize_html(cfg.kind+' · '+cfg.level),style),PageBreak()]
    for i,slide in enumerate(content['slides']):
        story.append(Paragraph(sanitize_html(slide.get('heading',f'Раздел {i+1}')),heading))
        for paragraph in slide.get('paragraphs',[])[:5]:story.append(Paragraph(sanitize_html(paragraph),style))
        if imgs[i]:
            im=PILImage.open(io.BytesIO(imgs[i]));w,h=im.size
            maxw,maxh=(640,285) if cfg.kind=='Презентация' else (450,310)
            ratio=min(maxw/w,maxh/h,1)
            story.append(Image(io.BytesIO(imgs[i]),width=w*ratio,height=h*ratio))
        if i<len(content['slides'])-1:story.append(PageBreak())
    if content.get('sources'):
        story.extend([PageBreak(),Paragraph('Источники',heading)])
        for src in content['sources'][:20]:story.append(Paragraph(sanitize_html(src),style))
    def decorate(canvas,document):
        canvas.saveState();w,h=page_size
        if cfg.background=='tinted':
            canvas.setFillColor(HexColor('#F4F0FF'));canvas.rect(0,0,w,h,stroke=0,fill=1)
        elif cfg.background=='dark':
            canvas.setFillColor(HexColor('#171326'));canvas.rect(0,0,w,h,stroke=0,fill=1)
        canvas.setStrokeColor(accent);canvas.setLineWidth(3);canvas.line(40,h-30,w-40,h-30)
        canvas.setFont('DocuUnicode',8);canvas.setFillColor(accent);canvas.drawRightString(w-42,22,str(document.page))
        canvas.restoreState()
    doc.build(story,onFirstPage=decorate,onLaterPages=decorate)

def add_docx(content,cfg,imgs,path):
    doc=Document();doc.add_heading(content.get('title',cfg.topic),0)
    for i,s in enumerate(content['slides']):
        doc.add_heading(s.get('heading',f'Раздел {i+1}'),level=1)
        for p in s.get('paragraphs',[]):doc.add_paragraph(p)
        if imgs[i]:doc.add_picture(io.BytesIO(imgs[i]),width=Inches(4.8))
    if content.get('sources'):
        doc.add_heading('Источники',1)
        for s in content['sources']:doc.add_paragraph(s)
    doc.save(path)

def add_pptx(content,cfg,imgs,path):
    prs=Presentation();prs.slide_width=PInches(13.33);prs.slide_height=PInches(7.5)
    for i,s in enumerate(content['slides']):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        bg=slide.background.fill;bg.solid()
        from pptx.dml.color import RGBColor
        bg.fore_color.rgb=RGBColor(22,18,40) if cfg.background=='dark' else RGBColor(248,246,255)
        color=RGBColor.from_string(COLORS.get(cfg.theme,COLORS['violet']).lstrip('#'))
        shape=slide.shapes.add_textbox(PInches(.8),PInches(.5),PInches(11.8),PInches(1.0))
        tf=shape.text_frame;tf.text=s.get('heading',f'Слайд {i+1}')
        tf.paragraphs[0].font.size=PPt(30);tf.paragraphs[0].font.bold=True;tf.paragraphs[0].font.color.rgb=color
        shape=slide.shapes.add_textbox(PInches(.9),PInches(1.7),PInches(6.4 if imgs[i] else 11),PInches(5))
        tf=shape.text_frame;tf.word_wrap=True
        for j,p in enumerate(s.get('paragraphs',[])[:3]):
            par=tf.paragraphs[0] if j==0 else tf.add_paragraph();par.text=p;par.font.size=PPt(18)
            par.font.color.rgb=RGBColor(240,235,255) if cfg.background=='dark' else RGBColor(39,34,60)
            par.space_after=PPt(18)
        if imgs[i]:slide.shapes.add_picture(io.BytesIO(imgs[i]),PInches(7.5),PInches(2),width=PInches(4.8))
    prs.save(path)

def make_formats(content,cfg,imgs,project,uid):
    files=[]
    for ext in dict.fromkeys(['pdf']+[x for x in cfg.formats if x in ('txt','json','md','docx','pptx')]):
        fid=secrets.token_hex(12);path=OUTPUT/f'{fid}.{ext}'
        if ext=='pdf':add_pdf(content,cfg,imgs,path)
        elif ext=='docx':add_docx(content,cfg,imgs,path)
        elif ext=='pptx':add_pptx(content,cfg,imgs,path)
        elif ext=='json':path.write_text(json.dumps({'config':cfg.model_dump(),'document':content},ensure_ascii=False,indent=2),encoding='utf-8')
        elif ext=='txt':path.write_text('\n\n'.join([content.get('title',cfg.topic)]+[s.get('heading','')+'\n'+'\n'.join(s.get('paragraphs',[])) for s in content['slides']]),encoding='utf-8')
        elif ext=='md':path.write_text('# '+content.get('title',cfg.topic)+'\n\n'+'\n\n'.join('## '+s.get('heading','')+'\n\n'+'\n\n'.join(s.get('paragraphs',[])) for s in content['slides']),encoding='utf-8')
        with conn() as db:db.execute('INSERT INTO files VALUES(?,?,?,?,?,?)',(fid,project,uid,ext,str(path),int(time.time())))
        files.append((fid,ext,path))
    return files

async def send_document(uid,path):
    if not BOT_TOKEN:return
    try:
        async with httpx.AsyncClient(timeout=75) as c:
            with path.open('rb') as f:
                r=await c.post(f'https://api.telegram.org/bot{BOT_TOKEN}/sendDocument',data={'chat_id':uid},files={'document':(path.name,f,'application/octet-stream')})
                if not r.json().get('ok'):logging.warning('sendDocument error: %s',r.text[:500])
    except Exception:logging.exception('Failed to send file to Telegram')

async def worker(job,uid,cfg):
    try:
        content=await generate_ai(cfg)
        imgs=[]
        for i,s in enumerate(content['slides']):
            raw=None
            if cfg.image_mode in ('custom','mixed') and i<len(cfg.images) and cfg.images[i]:
                with conn() as db:row=db.execute('SELECT path FROM uploads WHERE id=? AND user_id=?',(cfg.images[i],uid)).fetchone()
                if row:raw=Path(row['path']).read_bytes()
            if not raw and cfg.image_mode in ('auto','mixed'):raw=await image_auto(s.get('image_query',''))
            imgs.append(raw)
        project=secrets.token_hex(12)
        with conn() as db:db.execute('INSERT INTO projects VALUES(?,?,?,?,?,?,?)',(project,uid,content.get('title',cfg.topic),cfg.kind,cfg.model_dump_json(),json.dumps(content,ensure_ascii=False),int(time.time())))
        files=await asyncio.to_thread(make_formats,content,cfg,imgs,project,uid)
        with conn() as db:db.execute("UPDATE jobs SET status='done',project_id=? WHERE id=?",(project,job))
        for fid,ext,path in files:await send_document(uid,path)
    except Exception as exc:
        logging.exception('Job %s failed',job)
        with conn() as db:
            quota=db.execute('SELECT error FROM jobs WHERE id=?',(job,)).fetchone()['error']
            db.execute("UPDATE jobs SET status='error',error=? WHERE id=?",(str(exc)[:400],job))
            if quota in ('free','bonus'):
                col='free_left' if quota=='free' else 'bonus_left'
                db.execute(f'UPDATE users SET {col}={col}+1 WHERE id=?',(uid,))
        await tg_call('sendMessage',{'chat_id':uid,'text':'Не удалось создать документ. Откройте Mini App и повторите попытку.'})

@app.post('/api/generate')
async def generate(cfg:Config,uid=Depends(auth)):
    if cfg.kind not in ('Доклад','Презентация','Реферат','Проект','Конспект','Эссе','Курсовая','Статья','План','Сочинение'):raise HTTPException(400,'Неизвестный тип')
    if cfg.image_mode not in ('none','auto','custom','mixed'):raise HTTPException(400,'Неверный режим фото')
    if not API_KEY:raise HTTPException(503,'Администратор ещё не подключил AI API')
    jid=secrets.token_hex(12);now=int(time.time())
    with conn() as db:
        db.execute('BEGIN IMMEDIATE')
        u=db.execute('SELECT * FROM users WHERE id=?',(uid,)).fetchone()
        if u['premium_until']<=now:
            if u['bonus_left']>0:db.execute('UPDATE users SET bonus_left=bonus_left-1 WHERE id=?',(uid,));quota='bonus'
            elif u['free_left']>0:db.execute('UPDATE users SET free_left=free_left-1 WHERE id=?',(uid,));quota='free'
            else:raise HTTPException(402,'Бесплатные генерации закончились. Выберите Premium.')
        else:quota='premium'
        db.execute('INSERT INTO jobs VALUES(?,?,?,?,?,?)',(jid,uid,'processing',quota,None,now))
    asyncio.create_task(worker(jid,uid,cfg))
    return {'job_id':jid}

@app.get('/api/jobs/{jid}')
def job_status(jid:str,uid=Depends(auth)):
    with conn() as db:
        row=db.execute('SELECT * FROM jobs WHERE id=? AND user_id=?',(jid,uid)).fetchone()
        if not row:raise HTTPException(404)
        result=dict(row)
        if row['status']=='done':result['files']=[{'id':r['id'],'format':r['format']} for r in db.execute('SELECT * FROM files WHERE project_id=?',(row['project_id'],))]
    return result

@app.get('/api/projects')
def projects(uid=Depends(auth)):
    with conn() as db:
        rows=db.execute('SELECT id,title,kind,created,config FROM projects WHERE user_id=? ORDER BY created DESC LIMIT 100',(uid,)).fetchall()
        return [{'id':r['id'],'title':r['title'],'kind':r['kind'],'created':r['created'],'config':json.loads(r['config']),'files':[{'id':f['id'],'format':f['format']} for f in db.execute('SELECT id,format FROM files WHERE project_id=?',(r['id'],))]} for r in rows]

@app.get('/api/files/{fid}')
def download(fid:str,uid=Depends(auth)):
    with conn() as db:r=db.execute('SELECT * FROM files WHERE id=? AND user_id=?',(fid,uid)).fetchone()
    if not r:raise HTTPException(404)
    return FileResponse(r['path'],filename='DocuFlow.'+r['format'])

@app.delete('/api/projects/{pid}')
def delete_project(pid:str,uid=Depends(auth)):
    with conn() as db:
        row=db.execute('SELECT id FROM projects WHERE id=? AND user_id=?',(pid,uid)).fetchone()
        if not row:raise HTTPException(404)
        files=db.execute('SELECT path FROM files WHERE project_id=?',(pid,)).fetchall()
        db.execute('DELETE FROM files WHERE project_id=?',(pid,));db.execute('DELETE FROM projects WHERE id=?',(pid,))
    for f in files:Path(f['path']).unlink(missing_ok=True)
    return {'ok':True}
